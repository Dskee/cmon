"""把实现里新增的两处（分尺码库存、SIZE_STOCK_OUT）同步进方案文档。

只改三处：快照 schema 里的 stock 字段、紧跟其后的一段解释、事件表里的一行。
其余一字不动——这份文档是交稿版，改动越少越好核。
"""
from __future__ import annotations

import sys
from copy import deepcopy
from pathlib import Path

import docx
from docx.text.paragraph import Paragraph

SRC = Path(r"D:\AI\交付\贺梦阳-淘宝商品与竞品监控Agent方案-27届.docx")

STOCK_OLD_PREFIX = '"stock": { "tier"'
STOCK_NEW = (
    '"stock": { "tier": "in_stock | low | out", "raw_value": 0,',
    '"sizes": { "M": "in_stock", "L": "in_stock", "XL": "out" } },',
)

WHY = (
    "sizes 这一层是后加的。整款档位看不出局部缺货：一个尺码卖光、其余还有货，"
    "tier 仍然是 in_stock，而“降价 + 某尺码卖光 + 差评转向尺码”这三条同时出现，"
    "才是“对手在清这个尺码的库存”这个判断的证据。少了它，汇总层拿到两条互不相干的证据，"
    "只能得出“两者无关”。"
)

EVENT_ROW = (
    "分尺码缺货",
    "某一尺码跨到断货、整款仍有货",
    "局部缺货；整款断货仍归库存事件，不重复报",
)


def set_text(para: Paragraph, text: str) -> None:
    """保留首个 run 的字体，把整段文字换掉。"""
    runs = para.runs
    if not runs:
        para.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def main() -> int:
    doc = docx.Document(str(SRC))

    # 1) 快照 schema 的 stock 行
    target = next((p for p in doc.paragraphs if p.text.strip().startswith(STOCK_OLD_PREFIX)), None)
    if target is None:
        print("找不到 stock schema 段", file=sys.stderr)
        return 1
    old = target.text
    set_text(target, STOCK_NEW[0])
    tail = deepcopy(target._p)
    target._p.addnext(tail)
    set_text(Paragraph(tail, target._parent), STOCK_NEW[1])
    print("stock 行：", old.strip()[:60], "->", STOCK_NEW[0][:60])

    # 2) 解释段落：插在"三行中间值"那一段后面
    anchor = next((p for p in doc.paragraphs if "三行中间值" in p.text), None)
    if anchor is None:
        print("找不到 schema 解释段", file=sys.stderr)
        return 1
    if "sizes 这一层是后加的" in anchor.text:
        print("解释段已存在，跳过")
    else:
        new_p = deepcopy(anchor._p)
        anchor._p.addnext(new_p)
        set_text(Paragraph(new_p, anchor._parent), WHY)
        print("已插入 sizes 解释段")

    # 3) 事件表加一行
    table = next((t for t in doc.tables
                  if t.rows and t.rows[0].cells[0].text.strip() == "事件类型"), None)
    if table is None:
        print("找不到事件表", file=sys.stderr)
        return 1
    names = [r.cells[0].text.strip() for r in table.rows]
    if EVENT_ROW[0] in names:
        print("事件行已存在，跳过")
    else:
        table.add_row()
        for cell, text in zip(table.rows[-1].cells, EVENT_ROW, strict=False):
            cell.text = text
        print("事件表新增：", EVENT_ROW[0])
    print("事件表现在：", [r.cells[0].text.strip() for r in table.rows])

    doc.save(str(SRC))
    print("已保存", SRC)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())